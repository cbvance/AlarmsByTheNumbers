"""Master alarm database (MADB) workbook for rationalization meetings.

One row per alarm with the measured facts filled in and the
rationalization fields left blank for the team: cause, consequence,
corrective action, time to respond, severity, and the priority those
two produce from the site's priority matrix.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .badactors import diagnose
from .episodes import ParsedJournal
from .nuisance import (
    chatterers,
    find_chattering,
    find_fleeting,
    median_or_none,
)
from .site import PRIORITY_NAMES, Site, tag_of

RESPONSE_TIMES = [">30 min", "10 to 30 min", "3 to 10 min", "<3 min"]
SEVERITIES = ["Minor", "Major", "Severe"]
# Example matrix: rows severity, columns time to respond. A site's alarm
# philosophy owns the real one; edit the sheet, the formulas follow it.
MATRIX = [
    ["No alarm", "Low", "Low", "Medium"],
    ["Low", "Low", "Medium", "High"],
    ["Low", "Medium", "High", "High"],
]
CLASSES = [
    "Process",
    "Equipment",
    "Safety",
    "Environmental",
    "Diagnostic",
    "Not an alarm",
]

MEASURED = [
    "Tag",
    "Alarm",
    "Display Path",
    "Console",
    "As-Found Priority",
    "Activations",
    "Per Day",
    "Top-10 Rank",
    "Chattering",
    "Fleeting",
    "Median Active (min)",
    "Median Ack (s)",
    "First Guess",
]
TEAM = [
    "Classification",
    "Cause",
    "Consequence",
    "Corrective Action",
    "Time to Respond",
    "Severity",
    "Rationalized Priority",
    "Setpoint",
    "Deadband",
    "On Delay (s)",
    "Off Delay (s)",
    "Suppress In State",
    "Keep or Remove",
    "Rationalized By",
    "Date",
    "Notes",
]

HEAD = Font(bold=True, color="000000")
GREY = PatternFill("solid", fgColor="D9D9D9")
LIGHT = PatternFill("solid", fgColor="F2F2F2")
THIN = Border(*(Side(style="thin", color="808080"),) * 4)


# listing: export_madb
def export_madb(
    pj: ParsedJournal,
    path: str | Path,
    site: Site | None = None,
    source_name: str = "",
) -> Path:
    eps = pj.annunciated()
    by_src = defaultdict(list)
    for e in eps:
        by_src[e.source].append(e)
    chatter = find_chattering(eps)
    fleet = find_fleeting(eps, exclude=chatterers(chatter))
    rank = {
        s: i
        for i, (s, _) in enumerate(
            Counter(e.source for e in eps).most_common(10), 1
        )
    }

    rows = {}
    for src, lst in by_src.items():
        med = median_or_none([e.duration(pj.end) for e in lst])
        rows[src] = [
            tag_of(src),
            src.rsplit(":/alm:", 1)[-1],
            lst[0].displaypath,
            lst[0].console,
            PRIORITY_NAMES[lst[0].priority],
            len(lst),
            round(len(lst) / pj.days, 2),
            rank.get(src, ""),
            "Y" if src in chatter else "",
            fleet.get(src, "") or "",
            round(med / 60, 1) if med is not None else "",
            round(median_or_none([e.time_to_ack() for e in lst]) or 0, 0),
            diagnose(
                chatter[src].in_bursts / len(lst) if src in chatter else 0.0,
                fleet.get(src, 0),
                len(lst),
                med,
            ),
        ]
    if site:  # configured alarms that never fired still get rationalized
        for a in site.alarms:
            rows.setdefault(
                a.source,
                [
                    a.tag,
                    a.name,
                    a.displaypath,
                    a.console,
                    PRIORITY_NAMES[a.priority],
                    0,
                    0,
                    "",
                    "",
                    "",
                    "",
                    "",
                    "no activations in period",
                ],
            )

    wb = Workbook()
    ws = wb.active
    ws.title = "MADB"
    header = MEASURED + TEAM
    ws.append(header)
    for c, name in enumerate(header, 1):
        cell = ws.cell(1, c)
        cell.font, cell.border = HEAD, THIN
        cell.fill = GREY if c <= len(MEASURED) else LIGHT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ordered = sorted(rows.values(), key=lambda r: (-r[5], r[0], r[1]))
    tr = header.index("Time to Respond") + 1
    sv = header.index("Severity") + 1
    rp = header.index("Rationalized Priority") + 1
    for i, r in enumerate(ordered, start=2):
        ws.append(r + [""] * len(TEAM))
        t, s = get_column_letter(tr), get_column_letter(sv)
        ws.cell(i, rp).value = (
            f'=IF(OR({t}{i}="",{s}{i}=""),"",INDEX(Matrix!$B$2:$E$4,'
            f"MATCH({s}{i},Matrix!$A$2:$A$4,0),"
            f"MATCH({t}{i},Matrix!$B$1:$E$1,0)))"
        )
    last = len(ordered) + 1
    _validate(ws, tr, last, "Lookups!$A$2:$A$5")
    _validate(ws, sv, last, "Lookups!$B$2:$B$4")
    _validate(
        ws, header.index("Classification") + 1, last, "Lookups!$C$2:$C$7"
    )
    _validate(ws, header.index("Keep or Remove") + 1, last, '"Keep,Remove"')
    for c, name in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(c)].width = max(
            10, min(len(name) + 4, 34)
        )
    ws.column_dimensions["C"].width = 34
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{last}"

    m = wb.create_sheet("Matrix")
    m.append(["Severity \\ Time to Respond"] + RESPONSE_TIMES)
    for sev, row in zip(SEVERITIES, MATRIX):
        m.append([sev] + row)
    for row in m.iter_rows():
        for cell in row:
            cell.border = THIN
    for cell in m[1]:
        cell.font, cell.fill = HEAD, GREY
    m.column_dimensions["A"].width = 26
    for c in "BCDE":
        m.column_dimensions[c].width = 14

    lk = wb.create_sheet("Lookups")
    lk.append(["Time to Respond", "Severity", "Classification"])
    for i in range(max(len(RESPONSE_TIMES), len(SEVERITIES), len(CLASSES))):
        lk.append(
            [
                RESPONSE_TIMES[i] if i < len(RESPONSE_TIMES) else None,
                SEVERITIES[i] if i < len(SEVERITIES) else None,
                CLASSES[i] if i < len(CLASSES) else None,
            ]
        )

    info = wb.create_sheet("Period")
    for k, v in [
        ("Journal", source_name),
        ("From", pj.start),
        ("To", pj.end),
        ("Days", round(pj.days, 1)),
        ("Annunciated alarms", len(eps)),
        ("Alarms listed", len(ordered)),
        ("Exported", datetime.now().replace(microsecond=0)),
    ]:
        info.append([k, v])
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 40
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# end listing


def _validate(ws, col: int, last: int, formula: str) -> None:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    letter = get_column_letter(col)
    dv.add(f"{letter}2:{letter}{last}")


# listing: read_madb
@dataclass
class Decision:
    """One MADB row as the rationalization team left it."""

    row: int
    source: str
    displaypath: str
    first_guess: str
    chattering: bool
    team: dict[str, object]

    def get(self, field: str):
        v = self.team.get(field)
        return None if v in ("", None) else v


def read_madb(
    path: str | Path, provider: str = "default"
) -> tuple[list[Decision], dict]:
    """Read a filled MADB back, with the priority matrix it was filled in
    against. The workbook is the record; this turns it into data."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=False)
    ws = wb["MADB"]
    header = [c.value for c in ws[1]]
    col = {name: i for i, name in enumerate(header)}
    m = wb["Matrix"]
    times = [c.value for c in m[1]][1:]
    matrix = {
        (r[0].value, t): r[k + 1].value
        for r in m.iter_rows(min_row=2)
        for k, t in enumerate(times)
    }
    out = []
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not r[col["Tag"]]:
            continue
        team = {f: r[col[f]] for f in TEAM}
        rp = team["Rationalized Priority"]
        if isinstance(rp, str) and rp.startswith("="):
            team["Rationalized Priority"] = matrix.get(
                (team["Severity"], team["Time to Respond"])
            )
        out.append(
            Decision(
                i,
                f"prov:{provider}:/tag:{r[col['Tag']]}:/alm:{r[col['Alarm']]}",
                r[col["Display Path"]],
                r[col["First Guess"]] or "",
                r[col["Chattering"]] == "Y",
                team,
            )
        )
    return out, matrix


# end listing


# listing: check_madb
REQUIRED = [
    "Classification",
    "Cause",
    "Consequence",
    "Corrective Action",
    "Time to Respond",
    "Severity",
]


def check_madb(decisions: list[Decision]) -> list[str]:
    """Problems a reviewer would catch, found before the review."""
    problems = []
    for d in decisions:
        keep = d.get("Keep or Remove")
        if keep is None:
            continue  # not yet rationalized
        where = f"row {d.row} {d.displaypath}"
        if keep == "Remove":
            if not d.get("Notes"):
                problems.append(f"{where}: removed with no reason in Notes")
            continue
        for f in REQUIRED:
            if d.get(f) is None:
                problems.append(f"{where}: kept but {f} is blank")
        if d.get("Classification") == "Not an alarm":
            problems.append(f"{where}: classified Not an alarm but kept")
        if d.get("Rationalized Priority") == "No alarm":
            problems.append(f"{where}: matrix says No alarm but kept")
        if d.chattering and d.get("Deadband") is None:
            problems.append(f"{where}: chatters but has no deadband")
        if not d.get("Rationalized By"):
            problems.append(f"{where}: no name in Rationalized By")
    return problems


def progress(decisions: list[Decision]) -> dict[str, int]:
    """How far the team has got: rows decided, kept, removed, open."""
    keep = [d.get("Keep or Remove") for d in decisions]
    return {
        "alarms": len(decisions),
        "kept": keep.count("Keep"),
        "removed": keep.count("Remove"),
        "open": keep.count(None),
    }


# end listing


# The (severity, time to respond) answer that produces each priority on the
# example matrix, used to write a completed MADB from a site model.
ANSWER = {
    1: ("Major", "10 to 30 min"),
    2: ("Major", "3 to 10 min"),
    3: ("Severe", "3 to 10 min"),
}
PRIORITY_VALUE = {"Low": 1, "Medium": 2, "High": 3, "Critical": 4}


def _rationale(a) -> dict[str, str]:
    """Templated cause, consequence, and action for the synthetic plant.

    A real team writes these in the session; Red Mesa's are generated so the
    completed workbook can round-trip through the tools."""
    d, tag = a.description, a.tag.rsplit("/", 1)[-1]
    if a.name in ("H", "HH", "L", "LL"):
        way = "high" if a.name.startswith("H") else "low"
        return {
            "Cause": f"{d} {way}, past {a.setpoint:g} {a.units}",
            "Consequence": f"{d} out of range; equipment or product "
            f"at risk if it continues",
            "Corrective Action": f"Check {tag} trend and the loop "
            f"controlling it; correct the cause",
        }
    return {
        "Cause": f"{d}: {a.name.lower()}",
        "Consequence": "Loss of the equipment's function",
        "Corrective Action": f"Confirm {a.name.lower()} in the field; "
        f"start the spare or follow the procedure",
    }


# listing: madb_round_trip
def fill_from_site(
    path: str | Path,
    site: Site,
    by: str = "LM, CO, PS",
    date: str = "2026-07-14",
) -> Path:
    """Write the site model's rationalization decisions into an exported
    MADB, as the team would after its sessions."""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb["MADB"]
    header = [c.value for c in ws[1]]
    col = {name: i + 1 for i, name in enumerate(header)}
    by_disp = {a.displaypath: a for a in site.alarms}
    for r in range(2, ws.max_row + 1):
        a = by_disp.get(ws.cell(r, col["Display Path"]).value)
        if a is None:
            continue
        put = lambda f, v: ws.cell(r, col[f], v)  # noqa: E731
        put("Rationalized By", by)
        put("Date", date)
        if a.rat.remove:
            put("Classification", "Not an alarm")
            put("Keep or Remove", "Remove")
            put(
                "Notes",
                "Equipment status, not an abnormal condition "
                "needing a response; show on the graphic",
            )
            continue
        for f, v in _rationale(a).items():
            put(f, v)
        put(
            "Classification",
            (
                "Equipment"
                if a.name in ("Trip", "Fail", "Flame Fail")
                else "Process"
            ),
        )
        sev, time = ANSWER.get(a.rat.priority or a.priority, ANSWER[3])
        put("Severity", sev)
        put("Time to Respond", time)
        moved = a.rat.rate is not None
        step = 0.05 * abs(a.setpoint) or 1.0
        put(
            "Setpoint",
            (
                a.setpoint + (step if a.name.startswith("H") else -step)
                if moved
                else a.setpoint
            ),
        )
        for f, v in (
            ("Deadband", a.rat.deadband),
            ("On Delay (s)", a.rat.on_delay),
            ("Off Delay (s)", a.rat.off_delay),
        ):
            put(f, v or None)
        put("Suppress In State", ", ".join(a.rat.suppress_in) or None)
        put("Keep or Remove", "Keep")
    wb.save(path)
    return Path(path)


def site_from_madb(site: Site, decisions: list[Decision]) -> Site:
    """The site with every alarm's rationalized configuration taken from
    the MADB, so the generator applies what the team decided."""
    from dataclasses import replace

    from .site import Rationalized

    by_disp = {d.displaypath: d for d in decisions}
    alarms = []
    for a in site.alarms:
        d = by_disp.get(a.displaypath)
        if d is None or d.get("Keep or Remove") is None:
            alarms.append(a)
            continue
        sp = d.get("Setpoint")
        moved = sp is not None and float(sp) != a.setpoint
        rp = d.get("Rationalized Priority")
        states = d.get("Suppress In State") or ""
        rat = Rationalized(
            priority=PRIORITY_VALUE.get(rp) if rp else None,
            deadband=float(d.get("Deadband") or 0.0),
            on_delay=float(d.get("On Delay (s)") or 0.0),
            off_delay=float(d.get("Off Delay (s)") or 0.0),
            remove=d.get("Keep or Remove") == "Remove",
            rate=a.rat.rate if moved else None,
            suppress_in=tuple(
                s.strip() for s in states.split(",") if s.strip()
            ),
        )
        alarms.append(replace(a, rat=rat))
    return replace(site, alarms=alarms)


# end listing
