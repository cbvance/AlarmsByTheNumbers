from openpyxl import load_workbook

from almetrics.badactors import bad_actors
from almetrics.compare import compare
from almetrics.madb import export_madb
from almetrics.shelving import shelving_report
from almetrics.sites import redmesa
from conftest import SRC
from test_generator import pair  # noqa: F401  (module fixture)


def test_bad_actors_ranked(pair):  # noqa: F811
    before, _, _ = pair
    acts = bad_actors(before)
    assert len(acts) == 10
    assert [a.count for a in acts] == sorted(
        (a.count for a in acts), reverse=True
    )
    assert acts[0].diagnosis.startswith("chatter")
    assert abs(acts[-1].cum_pct - sum(a.pct for a in acts)) < 1e-9


def test_madb_workbook(pair, tmp_path):  # noqa: F811
    before, _, _ = pair
    p = export_madb(before, tmp_path / "madb.xlsx", redmesa.build(), "b.db")
    wb = load_workbook(p)
    assert wb.sheetnames == ["MADB", "Matrix", "Lookups", "Period"]
    ws = wb["MADB"]
    header = [c.value for c in ws[1]]
    assert "Consequence" in header and "Rationalized Priority" in header
    assert ws.max_row - 1 == len(
        redmesa.build().alarms
    )  # every configured alarm listed
    f = ws.cell(2, header.index("Rationalized Priority") + 1).value
    assert f.startswith("=IF(") and "Matrix!" in f


def test_shelving_and_compare(pair):  # noqa: F811
    before, after, _ = pair
    rb, ra = shelving_report(before), shelving_report(after)
    assert rb.shelved_activations > 0 and not rb.disabled_spans
    assert ra.state_events > 0
    rows = {r["metric"]: r for r in compare(before, after)}
    assert rows["Chattering alarms"]["after"] == 0
    assert rows["Annunciated per day"]["change_pct"] < -50


def test_what_if_and_cards(pair, tmp_path):  # noqa: F811
    from almetrics.badactors import bad_actors, what_if
    from almetrics.charts import actor_card

    before, _, _ = pair
    rows = what_if(before, top=5)
    assert [r["removed"] for r in rows] == [0, 1, 2, 3, 4, 5]
    assert rows[-1]["per_day"] < rows[0]["per_day"] / 2
    a = bad_actors(before, top=1)[0]
    assert actor_card(before, a, tmp_path / "card.png").exists()


def test_shelve_periods(builder):
    from almetrics.episodes import parse_journal
    from almetrics.shelving import shelve_periods

    for i, s in enumerate((0, 600, 1200, 9000, 9100)):
        builder.cycle(f"s{i}", SRC, s, s + 5, shelved=True)
    builder.cycle("x", SRC, 20000, 20005)
    ps = shelve_periods(parse_journal(builder.events()))
    assert [p.activations for p in ps] == [3, 2]
    assert ps[0].minutes == 20.0
